"""Generates the starter downloadable resources in public/resources/. Run: python3 scripts/build_resources.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.units import mm

OUT = Path(__file__).resolve().parent.parent / "public" / "resources"
OUT.mkdir(parents=True, exist_ok=True)
INK = "0D1B2A"; COPPER = "B4532A"; LINE = "E3E0D8"

# ---------------- RAID log (Excel) ----------------
wb = Workbook()
ins = wb.active; ins.title = "How to use"
rows = [
    ("RAID Log Template", None),
    ("by Rupesh Kumar · free to use and adapt", None),
    ("", None),
    ("Type", "R = Risk (might happen) · A = Assumption (believed true, needs validation) · I = Issue (has happened) · D = Dependency (needed from/by others)"),
    ("Writing risks", "Use cause → event → consequence. 'If X happens, then Y, because Z.'"),
    ("Scoring", "Probability and Impact on 1–5. Score = P × I. 15+ escalates to steering (red), 8–14 amber, below 8 green."),
    ("Dependencies", "Always record Date needed and Date committed. The gap is your early warning."),
    ("Cadence", "Review top-scored and overdue items weekly. Close items explicitly with a closure note."),
]
for r in rows: ins.append(list(r))
ins["A1"].font = Font(bold=True, size=16, color=INK)
ins["A2"].font = Font(italic=True, color="5A6472")
for r in range(4, 9):
    ins[f"A{r}"].font = Font(bold=True, color=INK)
    ins[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
ins.column_dimensions["A"].width = 18; ins.column_dimensions["B"].width = 100

ws = wb.create_sheet("RAID Log")
headers = ["ID", "Type", "Title", "Description (cause → event → consequence)", "Workstream", "Owner", "Raised",
           "Probability (1-5)", "Impact (1-5)", "Score", "RAG", "Mitigation / Action", "Date needed", "Date committed",
           "Due / Review", "Status", "Closure note"]
ws.append(headers)
thin = Side(style="thin", color=LINE)
for i, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=INK)
    c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(bottom=thin)
widths = [7, 7, 28, 48, 14, 14, 11, 11, 10, 8, 8, 40, 12, 12, 12, 11, 30]
for i, w in enumerate(widths, 1): ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
ws.row_dimensions[1].height = 34
examples = [
    ["R-001", "R", "Hardware delivery slip", "If the vendor misses the delivery date, wave 2 cutovers slip ≥3 weeks because change windows are booked 4 weeks ahead.", "Network", "PM", "2026-01-10", 3, 4, None, None, "Confirm dates in writing; identify alternative stock; pre-book backup windows.", None, None, "2026-01-24", "Open", ""],
    ["A-001", "A", "Out-of-hours site access", "We assume site access is available out of hours for all wave 1 sites.", "Facilities", "Site lead", "2026-01-10", 2, 3, None, None, "Validate with each site manager by end of month.", None, None, "2026-01-31", "Open", ""],
    ["D-001", "D", "Firewall rule approval", "Security team approval of new firewall rules required before cutover.", "Security", "Security lead", "2026-01-12", 3, 3, None, None, "Submit rules 3 weeks ahead; agree fast-track path.", "2026-02-05", "2026-02-10", "2026-02-01", "Open", ""],
    ["I-001", "I", "Test environment unavailable", "Lab environment down; regression testing for wave 1 delayed by 2 days.", "Network", "Eng lead", "2026-01-15", 5, 2, None, None, "Use vendor lab as interim; escalate repair.", None, None, "2026-01-17", "Open", ""],
]
for r in range(2, 202):
    ex = examples[r-2] if r-2 < len(examples) else [None]*len(headers)
    for ci, v in enumerate(ex, 1):
        if v is not None: ws.cell(row=r, column=ci, value=v)
    ws.cell(row=r, column=10, value=f'=IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r})),H{r}*I{r},"")')
    ws.cell(row=r, column=11, value=f'=IF(J{r}="","",IF(J{r}>=15,"Red",IF(J{r}>=8,"Amber","Green")))')
    for ci in range(1, len(headers)+1):
        ws.cell(row=r, column=ci).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=ci).border = Border(bottom=thin)
dv_type = DataValidation(type="list", formula1='"R,A,I,D"', allow_blank=True); ws.add_data_validation(dv_type); dv_type.add("B2:B201")
dv_score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws.add_data_validation(dv_score); dv_score.add("H2:I201")
dv_status = DataValidation(type="list", formula1='"Open,In progress,Escalated,Closed"', allow_blank=True); ws.add_data_validation(dv_status); dv_status.add("P2:P201")
ws.conditional_formatting.add("K2:K201", CellIsRule(operator="equal", formula=['"Red"'], fill=PatternFill("solid", fgColor="F8D7D3"), font=Font(color="8A1C12", bold=True)))
ws.conditional_formatting.add("K2:K201", CellIsRule(operator="equal", formula=['"Amber"'], fill=PatternFill("solid", fgColor="FBE9CF"), font=Font(color="8A5A00", bold=True)))
ws.conditional_formatting.add("K2:K201", CellIsRule(operator="equal", formula=['"Green"'], fill=PatternFill("solid", fgColor="DCEFE3"), font=Font(color="1E6B3A", bold=True)))
ws.freeze_panes = "D2"; ws.auto_filter.ref = "A1:Q201"
wb.save(OUT / "raid-log-template.xlsx")

# ---------------- PDF helpers ----------------
H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=HexColor("#"+INK), spaceAfter=4)
SUB = ParagraphStyle("sub", fontName="Helvetica", fontSize=9.5, leading=13, textColor=HexColor("#5A6472"), spaceAfter=14)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=HexColor("#"+COPPER), spaceBefore=12, spaceAfter=6)
BODY = ParagraphStyle("b", fontName="Helvetica", fontSize=9.8, leading=14, textColor=HexColor("#1B2A3D"))

def checklist_pdf(path, title, subtitle, sections):
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=18*mm, bottomMargin=16*mm,
                            title=title, author="Rupesh Kumar")
    story = [Paragraph(title, H1), Paragraph(subtitle, SUB)]
    for heading, items in sections:
        story.append(Paragraph(heading, H2))
        box = lambda: Table([[""]], colWidths=[3.4*mm], rowHeights=[3.4*mm], style=[("BOX", (0,0), (-1,-1), 0.8, HexColor("#"+COPPER))])
        data = [[box(), Paragraph(it, BODY)] for it in items]
        t = Table(data, colWidths=[8*mm, None])
        t.setStyle(TableStyle([("VALIGN", (0,0), (-1,-1), "TOP"), 
                               ("LINEBELOW", (0,0), (-1,-1), 0.4, HexColor("#"+LINE)), ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
        story.append(t)
    story += [Spacer(1, 14), Paragraph("Free to use and adapt. More resources and articles on technology, project management and careers on my website.", SUB)]
    doc.build(story)

checklist_pdf(OUT / "network-change-readiness-checklist.pdf",
    "Network Change Readiness Checklist",
    "Rupesh Kumar · A pre-cutover checklist for LAN/WAN/WLAN, firewall and data-centre changes",
    [
        ("1. Scope & approvals", [
            "Change record raised with clear scope, devices/sites in scope and business justification",
            "Change approved by CAB / change authority; maintenance window confirmed with business",
            "Stakeholders and service owners notified; communication plan agreed",
        ]),
        ("2. Design & validation", [
            "Low-level design / config reviewed and peer-checked",
            "Configuration tested in lab or on a representative pilot device",
            "Dependencies identified: circuits, firewall rules, DNS/DHCP, authentication, monitoring",
            "IP addressing, VLANs and routing changes documented",
        ]),
        ("3. Pre-change", [
            "Current configs backed up and stored outside the device",
            "Baseline captured: interface status, routing tables, key application tests",
            "Console / out-of-band access confirmed for every device",
            "Spares and vendor support (TAC) availability confirmed",
            "Monitoring alerts suppressed for the window (and re-enable step scheduled)",
        ]),
        ("4. Rollback", [
            "Rollback steps written and tested, with time required to roll back",
            "Go / no-go decision points and rollback trigger criteria agreed",
            "Named decision-maker available throughout the window",
        ]),
        ("5. Execution & validation", [
            "Runbook followed step by step with timestamps",
            "Post-change checks completed against the baseline",
            "Application owners / users confirm service (UAT) before closure",
        ]),
        ("6. Closure", [
            "Monitoring re-enabled; documentation and CMDB updated",
            "Hypercare period and contact defined",
            "Lessons learned captured for the next wave",
        ]),
    ])

checklist_pdf(OUT / "it-project-manager-interview-questions.pdf",
    "IT / Technical Project Manager Interview Question Bank",
    "Rupesh Kumar · Questions to practise for project, programme and technical PM interviews (infrastructure & IT)",
    [
        ("Delivery & planning", [
            "Walk me through how you plan a multi-site infrastructure rollout from kickoff to handover.",
            "How do you build and maintain a dependency map? Give an example where it saved a deadline.",
            "Tell me about a project that slipped. What did you see early, and what did you do?",
            "How do you decide between Agile, waterfall or hybrid for an infrastructure project?",
        ]),
        ("Risk, RAID & governance", [
            "How do you write a useful risk? Give an example from your experience.",
            "When and how do you escalate to a steering committee?",
            "Describe your weekly governance cadence and the reports you produce.",
            "How do you report a 'red' status to senior stakeholders without losing their confidence?",
        ]),
        ("Stakeholders & vendors", [
            "Tell me about a difficult stakeholder and how you handled them.",
            "How do you manage a vendor that is consistently late?",
            "How do you keep engineering teams motivated during a long programme?",
        ]),
        ("Technical understanding", [
            "Explain the difference between MPLS and SD-WAN to a non-technical stakeholder.",
            "What goes into a rollback plan for a firewall or core switch replacement?",
            "What are the main risks in a data-centre migration and how do you mitigate them?",
            "How would you approach a vulnerability-remediation programme across thousands of devices?",
        ]),
        ("Commercials & outcomes", [
            "How do you track budget versus forecast and explain variance?",
            "Tell me about a time you reduced cost or timeline. How did you measure it?",
            "How do you build a business case for a technology investment?",
        ]),
        ("Tips", [
            "Prepare 6–8 STAR stories covering delivery, risk, conflict, failure, leadership and cost.",
            "Quantify scope (sites, devices, users, budget) — and only use numbers you can defend.",
            "End each answer with the outcome and what you learned.",
        ]),
    ])
print("built:", [p.name for p in OUT.iterdir()])
